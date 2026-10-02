from io import BytesIO

import pandas as pd
from openpyxl import Workbook, load_workbook

from src.excel_writer import generate_distributed_excel_bytes


def _workbook_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "BASE"
    ws.append(["N° da nota", "Nota SGO", "Status do projeto", "P L N", "Projetistas"])
    ws.append([111, 430100001, "Em projeto", 5, None])
    ws.append([222, 430100002, "Em projeto", 7, "MARIA"])
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def test_generate_distributed_excel_bytes_keeps_source_and_adds_summary():
    source = _workbook_bytes()
    suggestions = pd.DataFrame([
        {
            "item_id": "excel-2",
            "Projetista": "JOAO",
            "Nº da nota": 111,
            "Nota SGO": 430100001,
            "PLN": 5,
            "Regional": "NORTE",
            "Município": "SAO LUIS",
            "Prazo": None,
            "Motivo": "Completar meta",
        }
    ])
    summary = pd.DataFrame([
        {
            "Projetista": "JOAO",
            "Novos projetos sugeridos": 1,
            "Novos postes sugeridos": 5,
            "Potencial postes após distribuição": 30,
            "Potencial projetos após distribuição": 5,
            "Risco de tempo": False,
            "Motivo": "Carteira cobre a meta diária",
        }
    ])

    output = generate_distributed_excel_bytes(source, suggestions, summary)

    source_wb = load_workbook(BytesIO(source), data_only=True)
    output_wb = load_workbook(BytesIO(output), data_only=True)
    assert source_wb["BASE"]["E2"].value is None
    assert output_wb["BASE"]["E2"].value == "JOAO"
    assert output_wb["BASE"]["E3"].value == "MARIA"
    assert "DISTRIBUICAO_AUTOMATICA" in output_wb.sheetnames
    summary_ws = output_wb["DISTRIBUICAO_AUTOMATICA"]
    assert summary_ws["B4"].value == 1
    assert summary_ws["A7"].value == "JOAO"
