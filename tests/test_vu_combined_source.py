from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

import pandas as pd
from openpyxl import Workbook, load_workbook

from src.distribution_engine import suggest_assignments
from src.excel_loader import load_vu_excel
from src.excel_writer import generate_vu_distributed_excel_bytes
from src.metrics import build_daily_snapshots
from src.models import StatusConfig, Targets


def make_vu_bytes():
    wb = Workbook()
    ws = wb.active
    ws.title = "VU"
    ws.append(["Projeto", "Ignorar", "Status", "Outra coluna"])
    ws.append(["440001", "x", "Em projeto", "não usar"])
    ws.append(["440002", "y", "Concluído", "não usar"])
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def test_vu_loader_reads_project_from_a_and_status_from_c_only():
    df = load_vu_excel(BytesIO(make_vu_bytes()))
    assert df.loc[0, "sgo"] == "440001"
    assert df.loc[0, "status"] == "Em projeto"
    assert df.loc[0, "posts"] == 0
    assert bool(df.loc[0, "posts_valid"]) is False
    assert df.loc[0, "assignee"] == ""
    assert df.loc[1, "sgo"] == "440002"
    assert df.loc[1, "status"] == "Concluído"


def test_vu_em_projeto_can_be_distributed_without_post_quantity():
    projects = load_vu_excel(BytesIO(make_vu_bytes()))
    projects = projects[projects["status"].eq("Em projeto")].copy()
    projects["source_base"] = "VU"
    projects["item_id"] = "VU::" + projects["item_id"].astype(str)

    targets = Targets(min_posts=25, target_posts=30, min_projects=4, target_projects=5)
    statuses = StatusConfig(project_pool=("Em projeto",), completed=("Concluído",))
    now = datetime(2026, 10, 2, 9, 0, tzinfo=ZoneInfo("America/Fortaleza"))
    snapshot = build_daily_snapshots(projects, ["Projetista A"], now, targets, statuses)

    suggestions, summary = suggest_assignments(
        projects,
        snapshot,
        targets,
        statuses,
        max_new_projects_per_designer=5,
        max_portfolio_posts=36,
        max_portfolio_projects=6,
        respect_time=False,
    )

    assert len(suggestions) == 1
    row = suggestions.iloc[0]
    assert row["Base de origem"] == "VU"
    assert row["Nota SGO"] == "440001"
    assert pd.isna(row["Postes Alterados/Novos"])
    assert row["PLN"] == 0
    assert row["Carga depois (projetos)"] == 1
    assert row["Carga depois (postes)"] == 0
    assert "conta somente para a meta de projetos" in row["Motivo"]


def test_vu_output_preserves_a_c_and_adds_assignee_column():
    suggestions = pd.DataFrame([
        {
            "item_id": "excel-vu-2",
            "Projetista": "Projetista A",
            "Base de origem": "VU",
            "Nota SGO": "440001",
            "Postes Alterados/Novos": None,
            "PLN": 0,
            "Carga antes (postes)": 0,
            "Carga depois (postes)": 0,
            "Carga antes (projetos)": 0,
            "Carga depois (projetos)": 1,
            "Motivo": "teste",
        }
    ])
    output = generate_vu_distributed_excel_bytes(make_vu_bytes(), suggestions)
    wb = load_workbook(BytesIO(output))
    ws = wb[wb.sheetnames[0]]
    assert ws["A2"].value == "440001"
    assert ws["C2"].value == "Em projeto"
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    assert "Projetista atribuído" in headers
    col = headers.index("Projetista atribuído") + 1
    assert ws.cell(2, col).value == "Projetista A"

from src.source_combiner import combine_levantamento_vu


def test_sources_do_not_combine_when_one_is_missing():
    vu = load_vu_excel(BytesIO(make_vu_bytes()))
    combined, overlap = combine_levantamento_vu(pd.DataFrame(), vu)
    assert combined.empty
    assert overlap == 0


def test_sources_complement_each_other_and_deduplicate_by_project_number():
    lev = pd.DataFrame([
        {"item_id": "LEVANTAMENTO::excel-2", "sgo": "440001", "status": "Em projeto", "posts": 5, "posts_valid": True, "source_base": "LEVANTAMENTO"},
        {"item_id": "LEVANTAMENTO::excel-3", "sgo": "430999", "status": "Concluído", "posts": 2, "posts_valid": True, "source_base": "LEVANTAMENTO"},
    ])
    vu = load_vu_excel(BytesIO(make_vu_bytes()))
    # Add one exclusive VU project in Em projeto.
    extra = vu.iloc[[0]].copy()
    extra["sgo"] = "440777"
    extra["item_id"] = "excel-vu-4"
    vu = pd.concat([vu, extra], ignore_index=True)
    vu["source_base"] = "VU"
    vu["item_id"] = "VU::" + vu["item_id"].astype(str)

    combined, overlap = combine_levantamento_vu(lev, vu)
    assert overlap == 1  # 440001 exists in both; keep LEVANTAMENTO only.
    assert set(combined["sgo"].astype(str)) == {"440001", "430999", "440777"}
    assert len(combined[combined["source_base"].eq("VU")]) == 1
