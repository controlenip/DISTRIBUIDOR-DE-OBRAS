from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

from src.excel_writer import apply_assignments_to_excel_copy


def test_writer_updates_only_copy(tmp_path: Path):
    source = tmp_path / "base.xlsx"
    output = tmp_path / "out.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["N° da nota", "Status do projeto ", "Projetistas"])
    ws.append([111, "Em projeto", None])
    wb.save(source)

    suggestions = pd.DataFrame([{"item_id": "excel-2", "Projetista": "JOAO"}])
    apply_assignments_to_excel_copy(source, output, suggestions)

    original = load_workbook(source, data_only=True).active
    generated = load_workbook(output, data_only=True).active
    assert original["C2"].value is None
    assert generated["C2"].value == "JOAO"
