from __future__ import annotations

from io import BytesIO
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def generate_daily_close_excel(snapshot: pd.DataFrame, analysis_date: date, target_posts: int, target_projects: int) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "FECHAMENTO_DIARIO"
    ws.sheet_view.showGridLines = False

    ws["A1"] = "FECHAMENTO DIÁRIO - PROJETISTAS"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="123B68")
    ws.merge_cells("A1:H1")
    ws["A2"] = "Data"
    ws["B2"] = analysis_date.strftime("%d/%m/%Y")
    ws["A3"] = "Meta diária"
    ws["B3"] = f"{target_posts} postes / {target_projects} projetos"

    cols = [
        "Projetista", "Postes realizados", "Projetos realizados", "Postes em carteira",
        "Projetos em carteira", "Previsão postes 18h", "Previsão projetos 18h", "Situação",
    ]
    row0 = 5
    for i, col in enumerate(cols, start=1):
        c = ws.cell(row0, i, col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2D6EA3")
        c.alignment = Alignment(horizontal="center")

    safe = snapshot.copy()
    for r_idx, (_, row) in enumerate(safe.iterrows(), start=row0 + 1):
        for c_idx, col in enumerate(cols, start=1):
            ws.cell(r_idx, c_idx, row.get(col))

    widths = [30, 17, 18, 17, 18, 20, 21, 26]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + idx)].width = width
    ws.freeze_panes = "A6"

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
