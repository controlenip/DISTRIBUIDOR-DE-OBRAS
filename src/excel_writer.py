from __future__ import annotations

from pathlib import Path
from shutil import copy2

import pandas as pd
from openpyxl import load_workbook

from .name_utils import normalize_column_label


def _find_header_column(ws, aliases: list[str]) -> int:
    lookup = {
        normalize_column_label(ws.cell(1, col).value): col
        for col in range(1, ws.max_column + 1)
        if ws.cell(1, col).value is not None
    }
    for alias in aliases:
        key = normalize_column_label(alias)
        if key in lookup:
            return lookup[key]
    raise ValueError(f"Coluna não encontrada no Excel. Esperado: {aliases}")


def apply_assignments_to_excel_copy(
    source_path: str | Path,
    output_path: str | Path,
    suggestions: pd.DataFrame,
) -> Path:
    """Apply suggested assignees to a COPY of BASE_LIST.xlsx.

    The original workbook is never edited. Only the Projetistas column is changed.
    Source rows come from excel_loader.load_base_excel().
    """
    source_path = Path(source_path)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    copy2(source_path, output_path)

    if suggestions.empty:
        return output_path

    wb = load_workbook(output_path)
    ws = wb[wb.sheetnames[0]]
    assignee_col = _find_header_column(ws, ["Projetistas", "Projetista"])
    status_col = _find_header_column(ws, ["Status do projeto", "Status do projeto "])

    for _, row in suggestions.iterrows():
        item_id = str(row.get("item_id", ""))
        if not item_id.startswith("excel-"):
            continue
        try:
            excel_row = int(item_id.split("-", 1)[1])
        except Exception as exc:
            raise ValueError(f"item_id Excel inválido: {item_id}") from exc

        current_status = str(ws.cell(excel_row, status_col).value or "").strip().casefold()
        current_assignee = str(ws.cell(excel_row, assignee_col).value or "").strip()
        if current_status != "em projeto":
            raise ValueError(f"Linha {excel_row}: status não é 'Em projeto'.")
        if current_assignee:
            raise ValueError(f"Linha {excel_row}: obra já possui projetista '{current_assignee}'.")

        ws.cell(excel_row, assignee_col).value = str(row["Projetista"])

    wb.save(output_path)
    return output_path
