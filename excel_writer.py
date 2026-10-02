from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path
from shutil import copy2

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .name_utils import normalize_column_label


def _find_header_column(ws, aliases: list[str]) -> int:
    lookup = {
        normalize_column_label(ws.cell(1, col).value): col
        for col in range(1, ws.max_column + 1)
        if ws.cell(1, col).value is not None
    }
    for alias in aliases:
        key = normalize_column_label(alias)
        if key in lookup:
            return lookup[key]
    raise ValueError(f"Coluna não encontrada no Excel. Esperado: {aliases}")


def _apply_assignments_to_workbook(wb, suggestions: pd.DataFrame) -> int:
    """Apply assignments in-memory and return the number of changed rows."""
    if suggestions.empty:
        return 0

    ws = wb[wb.sheetnames[0]]
    assignee_col = _find_header_column(ws, ["Projetistas", "Projetista"])
    status_col = _find_header_column(ws, ["Status do projeto", "Status do projeto "])

    applied = 0
    for _, row in suggestions.iterrows():
        item_id = str(row.get("item_id", ""))
        if not item_id.startswith("excel-"):
            continue
        try:
            excel_row = int(item_id.split("-", 1)[1])
        except Exception as exc:
            raise ValueError(f"item_id Excel inválido: {item_id}") from exc

        current_status = str(ws.cell(excel_row, status_col).value or "").strip().casefold()
        current_assignee = str(ws.cell(excel_row, assignee_col).value or "").strip()
        if current_status != "em projeto":
            raise ValueError(f"Linha {excel_row}: status não é 'Em projeto'.")
        if current_assignee:
            raise ValueError(f"Linha {excel_row}: obra já possui projetista '{current_assignee}'.")

        ws.cell(excel_row, assignee_col).value = str(row["Projetista"])
        applied += 1

    return applied


def _add_distribution_summary_sheet(
    wb,
    suggestions: pd.DataFrame,
    distribution_summary: pd.DataFrame | None = None,
    target_posts: int = 30,
    target_projects: int = 5,
    generated_at: datetime | None = None,
) -> None:
    """Add an audit-friendly summary without modifying the original data sheet layout."""
    sheet_name = "DISTRIBUICAO_AUTOMATICA"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    generated_at = generated_at or datetime.now()
    ws["A1"] = "DISTRIBUIÇÃO AUTOMÁTICA DE OBRAS"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    ws.merge_cells("A1:P1")
    ws["A2"] = "Gerado em"
    ws["B2"] = generated_at.strftime("%d/%m/%Y %H:%M:%S")
    ws["A3"] = "Meta diária"
    ws["B3"] = f"{target_posts} postes / {target_projects} projetos por projetista"
    ws["A4"] = "Novas obras distribuídas"
    ws["B4"] = int(len(suggestions))

    start_row = 6
    headers = [
        "Projetista", "Base de origem", "Nº da nota", "Nota SGO", "PI (Tipo Projeto)", "Dificuldade",
        "Experiência projetista", "Postes Alterados/Novos", "Prioridade", "Regional", "Município", "Prazo",
        "Carga antes", "Carga depois", "Projetos antes/depois", "Motivo"
    ]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(start_row, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r_idx, (_, row) in enumerate(suggestions.iterrows(), start=start_row + 1):
        values = [
            row.get("Projetista"),
            row.get("Base de origem"),
            row.get("Nº da nota"),
            row.get("Nota SGO"),
            row.get("PI (Tipo Projeto)"),
            row.get("Dificuldade"),
            row.get("Experiência projetista"),
            row.get("Postes Alterados/Novos", row.get("PLN")),
            row.get("Prioridade"),
            row.get("Regional"),
            row.get("Município"),
            row.get("Prazo"),
            f'{row.get("Carga antes (postes)", "")} postes',
            f'{row.get("Carga depois (postes)", "")} postes',
            f'{row.get("Carga antes (projetos)", "")} → {row.get("Carga depois (projetos)", "")}',
            row.get("Motivo"),
        ]
        for c_idx, value in enumerate(values, start=1):
            ws.cell(r_idx, c_idx, value)

    if distribution_summary is not None and not distribution_summary.empty:
        summary_start = start_row + len(suggestions) + 3
        ws.cell(summary_start, 1, "RESUMO POR PROJETISTA").font = Font(bold=True, size=12, color="FFFFFF")
        ws.cell(summary_start, 1).fill = PatternFill("solid", fgColor="1F4E78")
        ws.merge_cells(start_row=summary_start, start_column=1, end_row=summary_start, end_column=7)

        summary_cols = [
            "Projetista",
            "Bloqueado por projeto acima da meta",
            "Novos projetos sugeridos",
            "Novos postes sugeridos",
            "Potencial postes após distribuição",
            "Potencial projetos após distribuição",
            "Risco de tempo",
            "Motivo",
        ]
        for col, header in enumerate(summary_cols, start=1):
            cell = ws.cell(summary_start + 1, col, header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="5B9BD5")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for r_idx, (_, row) in enumerate(distribution_summary.iterrows(), start=summary_start + 2):
            for c_idx, col_name in enumerate(summary_cols, start=1):
                ws.cell(r_idx, c_idx, row.get(col_name))

    widths = {
        "A": 28, "B": 18, "C": 18, "D": 18, "E": 18, "F": 14, "G": 22, "H": 22,
        "I": 14, "J": 18, "K": 22, "L": 18, "M": 16, "N": 16, "O": 20, "P": 70
    }
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width
    ws.freeze_panes = "A7"
    ws.sheet_view.showGridLines = False


def apply_assignments_to_excel_copy(
    source_path: str | Path,
    output_path: str | Path,
    suggestions: pd.DataFrame,
) -> Path:
    """Apply suggested assignees to a COPY of BASE_LIST.xlsx.

    The original workbook is never edited. Only the Projetistas column is changed.
    Source rows come from excel_loader.load_base_excel().
    """
    source_path = Path(source_path)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    copy2(source_path, output_path)

    wb = load_workbook(output_path)
    _apply_assignments_to_workbook(wb, suggestions)
    wb.save(output_path)
    return output_path


def generate_distributed_excel_bytes(
    source_bytes: bytes,
    suggestions: pd.DataFrame,
    distribution_summary: pd.DataFrame | None = None,
    target_posts: int = 30,
    target_projects: int = 5,
    generated_at: datetime | None = None,
) -> bytes:
    """Return a new XLSX in memory with automatic assignments applied.

    Intended for Streamlit's download_button. The uploaded source bytes are never
    changed. Only the Projetistas column is written in the original data sheet;
    an additional summary sheet is appended for traceability.
    """
    if not source_bytes:
        raise ValueError("BASE LIST vazia ou não carregada.")

    wb = load_workbook(BytesIO(source_bytes))
    _apply_assignments_to_workbook(wb, suggestions)
    _add_distribution_summary_sheet(
        wb,
        suggestions,
        distribution_summary=distribution_summary,
        target_posts=target_posts,
        target_projects=target_projects,
        generated_at=generated_at,
    )

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def _apply_vu_assignments_to_workbook(wb, suggestions: pd.DataFrame) -> int:
    """Apply VU assignments without changing columns A/C.

    The VU input is read only from A (project number) and C (status). Because the
    source has no assignee field, the generated copy receives/uses a column named
    'Projetista atribuído' (created at D when absent). Columns A and C remain untouched.
    """
    if suggestions.empty:
        return 0
    ws = wb[wb.sheetnames[0]]
    # VU contract: A = project number, C = status.
    project_col = 1
    status_col = 3
    assignee_col = None
    for col in range(1, ws.max_column + 1):
        value = normalize_column_label(ws.cell(1, col).value)
        if value in {normalize_column_label("Projetista atribuído"), normalize_column_label("Projetistas"), normalize_column_label("Projetista")}:
            assignee_col = col
            break
    if assignee_col is None:
        assignee_col = max(4, ws.max_column + 1)
        ws.cell(1, assignee_col).value = "Projetista atribuído"
        ws.cell(1, assignee_col).font = Font(bold=True)

    applied = 0
    for _, row in suggestions.iterrows():
        item_id = str(row.get("item_id", ""))
        if not item_id.startswith("excel-vu-"):
            continue
        try:
            excel_row = int(item_id.rsplit("-", 1)[1])
        except Exception as exc:
            raise ValueError(f"item_id VU inválido: {item_id}") from exc

        current_status = str(ws.cell(excel_row, status_col).value or "").strip().casefold()
        current_assignee = str(ws.cell(excel_row, assignee_col).value or "").strip()
        if current_status != "em projeto":
            raise ValueError(f"Linha {excel_row}: status VU não é 'Em projeto'.")
        if current_assignee:
            raise ValueError(f"Linha {excel_row}: projeto VU já possui projetista '{current_assignee}'.")
        ws.cell(excel_row, assignee_col).value = str(row["Projetista"])
        applied += 1
    return applied


def generate_vu_distributed_excel_bytes(
    source_bytes: bytes,
    suggestions: pd.DataFrame,
    distribution_summary: pd.DataFrame | None = None,
    target_posts: int = 30,
    target_projects: int = 5,
    generated_at: datetime | None = None,
) -> bytes:
    """Generate a VU copy with assignments while preserving columns A and C."""
    if not source_bytes:
        raise ValueError("BASE LIST VU vazia ou não carregada.")
    wb = load_workbook(BytesIO(source_bytes))
    _apply_vu_assignments_to_workbook(wb, suggestions)
    _add_distribution_summary_sheet(
        wb,
        suggestions,
        distribution_summary=distribution_summary,
        target_posts=target_posts,
        target_projects=target_projects,
        generated_at=generated_at,
    )
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
